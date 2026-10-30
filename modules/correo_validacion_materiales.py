"""Generación de adjuntos y correo de Validación MO vs. Materiales."""

from __future__ import annotations

import html
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


# Destinatarios aprobados para Validación MO vs. Materiales.
# En modo revisión pueden verificarse o modificarse antes de enviar.
DESTINATARIOS = [
    "d.sanchez@eliteingenieros.com.co",
    "l.toro@eliteingenieros.com.co",
    "f.marin@eliteingenieros.com.co",
    "a.villegas@eliteingenieros.com.co",
    "y.patino@eliteingenieros.com.co",
    "y.betancur@eliteingenieros.com.co",
    "c.giraldo@eliteingenieros.com.co",
    "veronica.garzon@eliteingenieros.com.co",
    "dana.villa@eliteingenieros.com.co",
]
COPIAS = [
    "d.leon@eliteingenieros.com.co",
    "s.atencio@eliteingenieros.com.co",
]


def buscar_informe_mas_reciente(proyecto: Path) -> Path:
    """Devuelve el último informe generado por validar_export_almacen.py."""

    carpeta = proyecto / "outputs"
    archivos = list(carpeta.glob("validacion_mo_materiales_*.xlsx"))
    if not archivos:
        raise FileNotFoundError(
            "No se encontró un informe en la carpeta outputs. "
            "Ejecute primero 'Validar Materiales'."
        )
    return max(archivos, key=lambda archivo: archivo.stat().st_mtime)


def _nombre_seguro(nombre: str) -> str:
    nombre = re.sub(r"[^A-Za-z0-9_-]+", "_", nombre).strip("_")
    return nombre[:80] or "RESULTADO"


def generar_adjuntos_por_hoja(informe: Path) -> list[Path]:
    """
    Genera un libro por cada hoja de resultados.

    Cada adjunto conserva la hoja correspondiente y GUIA_REGLAS, incluyendo
    formato, filtros, tablas, colores, anchos y textos del informe original.
    """

    libro_origen = load_workbook(informe, read_only=True)
    hojas = [h for h in libro_origen.sheetnames if h != "GUIA_REGLAS"]
    libro_origen.close()

    comprobacion = load_workbook(informe, read_only=True)
    tiene_guia = "GUIA_REGLAS" in comprobacion.sheetnames
    comprobacion.close()
    if not tiene_guia:
        raise ValueError("El informe no contiene la hoja GUIA_REGLAS.")

    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    carpeta = informe.parent / "correos" / marca
    carpeta.mkdir(parents=True, exist_ok=True)

    adjuntos: list[Path] = []
    for numero, hoja_resultado in enumerate(hojas, start=1):
        libro = load_workbook(informe)
        conservar = {"GUIA_REGLAS", hoja_resultado}
        for hoja in list(libro.sheetnames):
            if hoja not in conservar:
                del libro[hoja]

        # La alerta queda visible al abrir; GUIA_REGLAS se conserva como apoyo.
        libro.active = libro.sheetnames.index(hoja_resultado)
        salida = carpeta / (
            f"{numero:02d}_{_nombre_seguro(hoja_resultado)}.xlsx"
        )
        libro.save(salida)
        libro.close()
        adjuntos.append(salida)

    return adjuntos


def construir_html(
    fecha_inicial: str,
    fecha_final: str,
    cantidad_adjuntos: int,
) -> str:
    """Construye un cuerpo HTML compacto y compatible con Outlook."""

    inicial = html.escape(fecha_inicial)
    final = html.escape(fecha_final)
    generado = datetime.now().strftime("%d/%m/%Y %I:%M %p")

    return f"""
    <html><body style="font-family:Segoe UI,Arial,sans-serif;color:#333;">
      <h2 style="color:#00995c;margin-bottom:8px;">
        Resultados de validación MO vs. Materiales
      </h2>
      <p>Se adjuntan los resultados de la validación realizada sobre la
      información extraída de Fénix para el período comprendido entre el
      <b>{inicial}</b> y el <b>{final}</b>.</p>
      <table style="border-collapse:collapse;margin:14px 0;">
        <tr><td style="padding:6px 12px;background:#eef7f1;"><b>Período</b></td>
            <td style="padding:6px 12px;">{inicial} al {final}</td></tr>
        <tr><td style="padding:6px 12px;background:#eef7f1;"><b>Estado</b></td>
            <td style="padding:6px 12px;color:#00995c;"><b>Completada</b></td></tr>
        <tr><td style="padding:6px 12px;background:#eef7f1;"><b>Adjuntos</b></td>
            <td style="padding:6px 12px;">{cantidad_adjuntos}</td></tr>
        <tr><td style="padding:6px 12px;background:#eef7f1;"><b>Generación</b></td>
            <td style="padding:6px 12px;">{generado}</td></tr>
      </table>
      <p>Cada archivo contiene una hoja de resultados y la hoja
      <b>GUIA_REGLAS</b> para facilitar su consulta.</p>
      <p style="background:#fff2cc;padding:10px;border-left:4px solid #f0ad00;">
      Las alertas orientan la revisión, pero deben ser validadas por el analista
      de acuerdo con la operación y las reglas de negocio vigentes.</p>
      <p style="margin:20px 0 14px 0;">Quedo atento.</p>

      <!-- Firma compatible con Outlook Desktop -->
      <table role="presentation" cellpadding="0" cellspacing="0" border="0"
             width="100%" style="width:100%;border-collapse:collapse;
             border-top:1px solid #d7dee5;">
        <tr>
          <td width="7" style="width:7px;background:#008c7a;"></td>
          <td style="padding:14px 16px;font-family:Segoe UI,Arial,sans-serif;">
            <div style="font-size:14px;font-weight:bold;color:#153b3d;">
              Héctor Alejandro Gaviria Marín
            </div>
            <div style="font-size:11px;font-weight:bold;color:#008c7a;
                        margin-top:5px;">
              Analista de Datos
            </div>
            <div style="font-size:10px;color:#53677a;margin-top:4px;">
              Automatización y Analítica Operativa
            </div>
          </td>
        </tr>
      </table>

      <table role="presentation" cellpadding="0" cellspacing="0" border="0"
             width="100%" style="width:100%;border-collapse:collapse;">
        <tr>
          <td align="center" style="background:#1f2b3a;color:#ffffff;
              padding:13px 10px;font-family:Segoe UI,Arial,sans-serif;
              font-size:9px;">
            Informe generado automáticamente por
            <b> VALIDACIÓN MO VS. MATERIALES</b>.
          </td>
        </tr>
      </table>
    </body></html>
    """


def crear_correo_outlook(
    fecha_inicial: str,
    fecha_final: str,
    adjuntos: list[Path],
    enviar_automaticamente: bool = False,
) -> None:
    """Abre el correo para revisión o lo envía mediante Outlook Desktop."""

    import win32com.client as win32

    outlook = win32.Dispatch("Outlook.Application")
    correo = outlook.CreateItem(0)
    correo.To = ";".join(DESTINATARIOS)
    correo.CC = ";".join(COPIAS)
    correo.Subject = (
        "[Validación MO vs. Materiales] Resultados del período "
        f"{fecha_inicial} al {fecha_final}"
    )
    correo.Importance = 1
    correo.HTMLBody = construir_html(
        fecha_inicial, fecha_final, len(adjuntos)
    )

    for adjunto in adjuntos:
        correo.Attachments.Add(str(adjunto.resolve()))

    if enviar_automaticamente:
        correo.Send()
    else:
        correo.Display()

